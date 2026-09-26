#!/usr/bin/env python3
"""Extract a PowerPoint deck into Slide Forge inputs.

    python scripts/extract_pptx.py talk.pptx -o work/

Writes:
  work/outline.md      draft outline (one "## [L-…] headline" per slide, bullets, notes)
  work/content.json    full text / notes / image list per slide
  work/img/…           embedded pictures (reference them as img/<file> in deck.src.html)

The draft keeps the original wording. Headlines that are topic labels ("背景", "まとめ")
will be flagged by `outline.py check` — rewrite them as claims before scaffolding.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    sys.exit("✗ python-pptx missing:  pip install python-pptx")


def texts(shape):
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip()
            if t:
                yield p.level, t
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            yield from texts(s)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            yield 0, " | ".join(c.text.strip() for c in row.cells)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx", type=Path)
    ap.add_argument("-o", "--out", type=Path, default=Path("work"))
    a = ap.parse_args()
    prs = Presentation(str(a.pptx))
    (a.out / "img").mkdir(parents=True, exist_ok=True)
    slides, md = [], [f"# {a.pptx.stem}", "- density: speaker", "- audience: ", "- goal: ", ""]
    for i, sl in enumerate(prs.slides, 1):
        title = sl.shapes.title.text.strip() if sl.shapes.title is not None and sl.shapes.title.has_text_frame else ""
        body, imgs = [], []
        for sh in sl.shapes:
            if sl.shapes.title is not None and sh.shape_id == sl.shapes.title.shape_id:
                continue
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                name = f"s{i:02d}-{len(imgs) + 1}.{sh.image.ext}"
                (a.out / "img" / name).write_bytes(sh.image.blob)
                imgs.append(name)
            body += [(lv, t) for lv, t in texts(sh)]
        notes = sl.notes_slide.notes_text_frame.text.strip() if sl.has_notes_slide and sl.notes_slide.notes_text_frame else ""
        slides.append({"index": i, "title": title, "body": [t for _, t in body], "images": imgs, "notes": notes})
        layout = "L-title" if i == 1 else ("L-image" if imgs and len(body) <= 1 else ("L-split" if imgs else "L-bullets"))
        md.append(f"## [{layout}] {title or (body[0][1] if body else f'スライド {i}')}")
        for lv, t in body[:8]:
            md.append(("  " * lv) + f"- {t}")
        if imgs:
            md.append(f"images: {', '.join('img/' + n for n in imgs)}")
        if notes:
            md.append("notes: " + notes.replace("\n", " "))
        md.append("")
    if slides and slides[0]["title"]:
        md[0] = f"# {slides[0]['title']}"
    (a.out / "content.json").write_text(json.dumps(slides, ensure_ascii=False, indent=1), encoding="utf-8")
    (a.out / "outline.md").write_text("\n".join(md), encoding="utf-8")
    print(f"✓ {len(slides)} slides → {a.out}/outline.md, content.json, img/ ({sum(len(s['images']) for s in slides)} images)")
    print("  next: fill audience/goal, rewrite headlines as claims, then outline.py check")


if __name__ == "__main__":
    main()
